"""Store costs: what the business pays the landlord for each store, and the importer for
the owner's template (docs/qa/PLAN-010b-setup-fees-store-costs-pnl.md D8, D10, D12).

Rules kept here, not in the views or the command:

- **Date-effective, no overlap by construction.** :func:`set_cost` locks the store row,
  then compares with the store's open cost row (at most one, a partial unique index):
  the same values change nothing; the same ``effective_from`` updates the open row in
  place; an earlier ``effective_from`` is refused with a sentence; a later one closes
  the open row on ``effective_from`` minus one day and adds a new open row, in the one
  transaction. ``effective_to`` is the last day a row is in effect.
- **Unknown is never zero.** :func:`current_cost` is the row in effect on a day, or
  ``None``. :func:`margin` is the operator rent of the lease active that day (from
  ``lease_rent_history``) minus the landlord rent and the other monthly cost; ``None``
  when there is no cost row, no active lease or no rent row that day. An empty other
  cost on a cost row means there is none.
- **The importer** (:func:`parse_file`, :func:`import_rows`, the ``import_store_costs``
  command) reads the template's ``.csv`` (UTF-8, with or without a byte order mark) or
  ``.xlsx`` (sheet "Store costs" by name, read-only, cached values only: a formula is
  never evaluated, and one without a saved value refuses its row). A file that cannot be
  read, has more than :data:`MAX_ROWS` rows or lacks a required column is refused as a
  whole, in one sentence. Each row is then applied, or refused with its row number and
  a reason that never repeats a cell's value, in its own transaction: a refused row
  writes nothing, and the other rows are applied. Every cell of every row goes through
  the card-number guard. ``dry_run`` reads the open cost rows and writes nothing.

- **Corrections** (PLAN-010c D10, admin, a note required): :func:`correct_cost` writes
  a new row with the corrected values and dates and marks the old one superseded, in
  one transaction under the store's row lock; :func:`close_cost` ends the open row on a
  date. Only *live* rows (``superseded_at`` null) count anywhere: the open row, the
  overlap rule, :func:`current_cost`, :func:`margin`, the importer's comparison and the
  P&L. :func:`history` lists every row, live ones first. A result that would overlap
  another live row, or end before it starts, is refused with a sentence. Months after a
  closed row have no cost row and read "unknown"; a row with landlord rent 0.00 is a
  known zero.

Every write is audited: ``audit_log`` (entity ``store_cost``) for every change, plus one
``import_audit`` row (``target_table="store_costs"``) per row the importer applied.
Money is ``Decimal`` throughout. ``openpyxl`` is imported here only, and only when an
``.xlsx`` file is read.
"""

from __future__ import annotations

import csv
import io
import numbers
import re
import warnings
import zipfile
from collections import Counter
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from django.contrib.auth.models import User
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.db.models import QuerySet
from django.utils import timezone

from accounts.roles import Role, meets, role_of
from ledger.models import (
    NAME_MAX_LENGTH,
    AuditLog,
    ImportAudit,
    Lease,
    Store,
    StoreCost,
    StoreCostSource,
)
from ledger.services import leases
from ledger.services.rent_run import CENT, MONEY_MAX
from ledger.validators import PAYMENT_DATA_MESSAGE, validate_no_payment_data, visible_text

ENTITY_STORE_COST = "store_cost"
TARGET_TABLE = "store_costs"

# --- Fixed sentences ------------------------------------------------------------------

RENT_NEGATIVE = "The landlord rent cannot be negative."
OTHER_NEGATIVE = "The other monthly cost cannot be negative."
NOT_CENTS = "Amounts must have at most two decimals."
TOO_LARGE = "The amount is too large."
DESCRIPTION_TOO_LONG = f"The description may have at most {NAME_MAX_LENGTH} characters."
EARLIER_THAN_OPEN = (
    "The effective date is before {start}, when the current landlord rent started. "
    "Nothing was saved."
)
AFTER_CLOSED = (
    "The effective date must be after {end}, when the last landlord rent row ends. "
    "Nothing was saved."
)
SAVED_NEW = "Landlord rent saved from {start}."
SAVED_IN_PLACE = "Landlord rent from {start} updated."
UNCHANGED = "Nothing was changed."

ACTION_CREATED = "created"
ACTION_UPDATED = "updated"
ACTION_UNCHANGED = "unchanged"
ACTION_REFUSED = "refused"

# --- Outcomes -------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class CostOutcome:
    ok: bool
    #: One of ``created``, ``updated``, ``unchanged``, ``refused``.
    action: str
    message: str
    #: The form field a refusal belongs to (``None``: the whole form).
    field: str | None = None
    cost: StoreCost | None = None
    #: The open row this one closed, when it closed one.
    closed: StoreCost | None = None


@dataclass(frozen=True, slots=True)
class CostValues:
    """The values of one cost row, as staff typed them or as a file row reads."""

    landlord_rent: Decimal
    other_monthly_cost: Decimal | None
    other_cost_description: str
    landlord_lease_end: date | None
    notes: str

    def same_as(self, row: StoreCost) -> bool:
        return (
            self.landlord_rent,
            self.other_monthly_cost,
            self.other_cost_description,
            self.landlord_lease_end,
            self.notes,
        ) == (
            row.landlord_rent,
            row.other_monthly_cost,
            row.other_cost_description,
            row.landlord_lease_end,
            row.notes,
        )


# --- Reads ----------------------------------------------------------------------------


def live_costs(store: Store) -> QuerySet[StoreCost]:
    """The store's cost rows that count: none superseded (D10)."""
    return StoreCost.objects.filter(store=store, superseded_at__isnull=True)


def open_cost(store: Store) -> StoreCost | None:
    """The store's live row with no end (one at most, by constraint)."""
    return live_costs(store).filter(effective_to__isnull=True).first()


def latest_cost(store: Store) -> StoreCost | None:
    """The store's live row that starts last: the open row when there is one (live rows
    never overlap), else the row that ends last."""
    return live_costs(store).order_by("-effective_from").first()


def current_cost(store: Store, on: date) -> StoreCost | None:
    """The live cost row in effect on ``on``, or ``None``: the cost is unknown that day."""
    return (
        live_costs(store)
        .filter(effective_from__lte=on)
        .exclude(effective_to__lt=on)
        .order_by("-effective_from")
        .first()
    )


def history(store: Store) -> list[StoreCost]:
    """Every cost row of the store: the live ones first, then the superseded ones, each
    part latest first."""
    rows = StoreCost.objects.filter(store=store).order_by("-effective_from", "-created_at")
    return sorted(rows, key=lambda row: row.superseded_at is not None)


@dataclass(frozen=True, slots=True)
class HistoryEntry:
    """One row of the store page's cost history, with who superseded or closed it."""

    row: StoreCost
    #: The username of the admin whose correction superseded the row, or "".
    superseded_by_user: str = ""
    #: The username of the admin who closed the row, or "".
    closed_by_user: str = ""


def history_entries(store: Store) -> list[HistoryEntry]:
    """:func:`history` with the admin of each correction and close, read from the audit
    log (2 queries)."""
    rows = history(store)
    actors: dict[tuple[str, str], str] = {}
    for action, entity_id, username in (
        AuditLog.objects.filter(
            entity_type=ENTITY_STORE_COST,
            action__in=(ACTION_AUDIT_SUPERSEDED, ACTION_AUDIT_CLOSED_BY_ADMIN),
            entity_id__in=[str(row.pk) for row in rows],
        )
        .order_by("at")
        .values_list("action", "entity_id", "actor__username")
    ):
        actors[(action, entity_id)] = username or ""
    return [
        HistoryEntry(
            row=row,
            superseded_by_user=actors.get((ACTION_AUDIT_SUPERSEDED, str(row.pk)), ""),
            closed_by_user=actors.get((ACTION_AUDIT_CLOSED_BY_ADMIN, str(row.pk)), ""),
        )
        for row in rows
    ]


def active_lease(store: Store, on: date) -> Lease | None:
    """The store's lease active on ``on``: started on or before it and not ended before
    it. On a hand-over day (one lease ends, the next starts) the later start wins."""
    return (
        Lease.objects.filter(store=store, started_on__lte=on)
        .exclude(ended_on__lt=on)
        .order_by("-started_on", "-created_at")
        .first()
    )


def operator_rent(store: Store, on: date) -> Decimal | None:
    """The rent of the lease active on ``on``, or ``None`` when unknown."""
    lease = active_lease(store, on)
    if lease is None:
        return None
    row = leases.rent_on(lease, on)
    return Decimal(row.amount) if row is not None else None


def margin(store: Store, on: date) -> Decimal | None:
    """Operator rent minus landlord rent minus other monthly cost on ``on``.

    ``None`` when any input is unknown (no cost row, no active lease, no rent row);
    never zero for an unknown. An empty other cost on the cost row means none.
    """
    cost = current_cost(store, on)
    rent = operator_rent(store, on)
    if cost is None or rent is None:
        return None
    other = Decimal(cost.other_monthly_cost) if cost.other_monthly_cost is not None else Decimal(0)
    return (rent - Decimal(cost.landlord_rent) - other).quantize(CENT)


# --- Writes ---------------------------------------------------------------------------


def _money_refusal(value: Decimal | None, *, field: str, message: str) -> CostOutcome | None:
    if value is None:
        return None
    if not isinstance(value, Decimal) or not value.is_finite() or value < 0:
        return CostOutcome(False, ACTION_REFUSED, message, field)
    if value != value.quantize(CENT):
        return CostOutcome(False, ACTION_REFUSED, NOT_CENTS, field)
    if value > MONEY_MAX:
        return CostOutcome(False, ACTION_REFUSED, TOO_LARGE, field)
    return None


def _text_refusal(text: str, field: str) -> CostOutcome | None:
    try:
        validate_no_payment_data(text)
    except ValidationError:
        return CostOutcome(False, ACTION_REFUSED, PAYMENT_DATA_MESSAGE, field)
    return None


def values_refusal(values: CostValues) -> CostOutcome | None:
    """The first rule ``values`` break before anything is read, or ``None``."""
    if len(values.other_cost_description) > NAME_MAX_LENGTH:
        return CostOutcome(False, ACTION_REFUSED, DESCRIPTION_TOO_LONG, "other_cost_description")
    return (
        _money_refusal(values.landlord_rent, field="landlord_rent", message=RENT_NEGATIVE)
        or _money_refusal(
            values.other_monthly_cost, field="other_monthly_cost", message=OTHER_NEGATIVE
        )
        or _text_refusal(values.other_cost_description, "other_cost_description")
        or _text_refusal(values.notes, "notes")
    )


def decide(latest: StoreCost | None, values: CostValues, effective_from: date) -> CostOutcome:
    """What :func:`set_cost` would do against the store's latest live row
    (:func:`latest_cost`); reads and writes nothing.

    With no open row (none at all, or the latest was closed) a new open row is created,
    and it must start after the latest row ends, so live rows never overlap. With an open
    row: unchanged, updated in place, refused, or a new row that closes it. The result
    carries no row; the dry run of the importer reports it as it is.
    """
    if latest is not None and latest.effective_to is not None:
        if effective_from <= latest.effective_to:
            return CostOutcome(
                False,
                ACTION_REFUSED,
                AFTER_CLOSED.format(end=latest.effective_to.isoformat()),
                "effective_from",
            )
        return CostOutcome(True, ACTION_CREATED, SAVED_NEW.format(start=effective_from))
    open_row = latest
    if open_row is None:
        return CostOutcome(True, ACTION_CREATED, SAVED_NEW.format(start=effective_from))
    if values.same_as(open_row):
        return CostOutcome(True, ACTION_UNCHANGED, UNCHANGED)
    if effective_from < open_row.effective_from:
        return CostOutcome(
            False,
            ACTION_REFUSED,
            EARLIER_THAN_OPEN.format(start=open_row.effective_from.isoformat()),
            "effective_from",
        )
    if effective_from == open_row.effective_from:
        return CostOutcome(True, ACTION_UPDATED, SAVED_IN_PLACE.format(start=effective_from))
    return CostOutcome(True, ACTION_CREATED, SAVED_NEW.format(start=effective_from))


def set_cost(
    store: Store,
    landlord_rent: Decimal,
    effective_from: date,
    other_monthly_cost: Decimal | None = None,
    other_cost_description: str = "",
    landlord_lease_end: date | None = None,
    notes: str = "",
    source: str = StoreCostSource.STAFF,
    actor: User | None = None,
) -> CostOutcome:
    """Set the store's landlord rent from ``effective_from`` (D8), under the store's row
    lock: unchanged, updated in place, refused, or a new open row (see the module)."""
    values = CostValues(
        landlord_rent=landlord_rent,
        other_monthly_cost=other_monthly_cost,
        other_cost_description=visible_text(other_cost_description or ""),
        landlord_lease_end=landlord_lease_end,
        notes=visible_text(notes or ""),
    )
    refused = values_refusal(values)
    if refused is not None:
        return refused
    try:
        with transaction.atomic():
            locked = Store.objects.select_for_update().get(pk=store.pk)
            latest = latest_cost(locked)
            outcome = decide(latest, values, effective_from)
            if not outcome.ok or outcome.action == ACTION_UNCHANGED:
                return outcome
            open_row = latest if latest is not None and latest.effective_to is None else None
            before = _state(open_row) if open_row is not None else None
            closed: StoreCost | None = None
            if outcome.action == ACTION_UPDATED and open_row is not None:
                row = open_row
                _assign(row, values, source)
                row.save()
                _audit(actor, "store_cost.updated", row, before, _state(row))
                return CostOutcome(True, ACTION_UPDATED, outcome.message, cost=row)
            if open_row is not None:
                open_row.effective_to = effective_from - timedelta(days=1)
                open_row.save(update_fields=["effective_to", "updated_at"])
                _audit(actor, "store_cost.closed", open_row, before, _state(open_row))
                closed = open_row
            row = StoreCost(store=locked, effective_from=effective_from)
            _assign(row, values, source)
            row.save()
            _audit(actor, "store_cost.created", row, None, _state(row))
            return CostOutcome(True, ACTION_CREATED, outcome.message, cost=row, closed=closed)
    except ValidationError:
        # The card-number guard on save; the transaction is rolled back.
        return CostOutcome(False, ACTION_REFUSED, PAYMENT_DATA_MESSAGE)
    except IntegrityError:
        # Not reachable under the store lock; kept so a constraint never becomes a 500.
        return CostOutcome(
            False, ACTION_REFUSED, "The landlord rent could not be saved. Nothing was saved."
        )


def _assign(row: StoreCost, values: CostValues, source: str) -> None:
    row.landlord_rent = values.landlord_rent
    row.other_monthly_cost = values.other_monthly_cost
    row.other_cost_description = values.other_cost_description
    row.landlord_lease_end = values.landlord_lease_end
    row.notes = values.notes
    row.source = source


def _state(row: StoreCost) -> dict[str, Any]:
    """The audited fields of a cost row."""
    return {
        "store": str(row.store_id),
        "landlord_rent": str(row.landlord_rent),
        "other_monthly_cost": (
            str(row.other_monthly_cost) if row.other_monthly_cost is not None else None
        ),
        "other_cost_description": row.other_cost_description,
        "landlord_lease_end": (
            row.landlord_lease_end.isoformat() if row.landlord_lease_end else None
        ),
        "effective_from": row.effective_from.isoformat(),
        "effective_to": row.effective_to.isoformat() if row.effective_to else None,
        "source": row.source,
        "notes": row.notes,
        "superseded_at": row.superseded_at.isoformat() if row.superseded_at else None,
        "superseded_by": str(row.superseded_by_id) if row.superseded_by_id else None,
        "change_note": row.change_note,
    }


def _audit(
    actor: User | None,
    action: str,
    row: StoreCost,
    before: dict[str, Any] | None,
    after: dict[str, Any] | None,
) -> None:
    AuditLog.objects.create(
        actor=actor,
        action=action,
        entity_type=ENTITY_STORE_COST,
        entity_id=str(row.pk),
        before=before,
        after=after,
    )


# --- Corrections and closing (PLAN-010c D10) ------------------------------------------

CHANGE_NOTE_MAX_LENGTH = 2000
CHANGE_NOTE_REQUIRED = "A note is required: say why the row is changed."
CLOSE_NOTE_REQUIRED = "A note is required: say why the row is closed."
CHANGE_NOTE_TOO_LONG = f"The note may have at most {CHANGE_NOTE_MAX_LENGTH} characters."
END_BEFORE_START = "The end date cannot be before the start date. Nothing was saved."
ROW_NOT_FOUND = "This cost row is not one of this store's. Nothing was saved."
ALREADY_SUPERSEDED = "This cost row was already corrected. Nothing was saved."
OVERLAPS = (
    "The corrected dates overlap the row from {start} to {end}. "
    "Correct or close that row first. Nothing was saved."
)
NO_OPEN_ROW = "The store has no open landlord rent row to close. Nothing was saved."
CLOSE_BEFORE_START = (
    "The end date is before {start}, when this landlord rent started. Nothing was saved."
)
CORRECTED = "Cost row corrected; the old row is kept in the history, struck through."
CLOSED = "Landlord rent closed on {end}. Later months have no cost row and read unknown."

ACTION_AUDIT_SUPERSEDED = "store_cost.superseded"
ACTION_AUDIT_CORRECTION = "store_cost.correction_created"
ACTION_AUDIT_CLOSED_BY_ADMIN = "store_cost.closed_by_admin"


def _require_admin(actor: User | None) -> None:
    if actor is None or not meets(role_of(actor), Role.ADMIN):
        raise PermissionDenied


def _note_refusal(note: str, required: str) -> CostOutcome | None:
    if not note:
        return CostOutcome(False, ACTION_REFUSED, required, "change_note")
    if len(note) > CHANGE_NOTE_MAX_LENGTH:
        return CostOutcome(False, ACTION_REFUSED, CHANGE_NOTE_TOO_LONG, "change_note")
    return _text_refusal(note, "change_note")


def _overlapping(
    store: Store, start: date, end: date | None, exclude: StoreCost
) -> StoreCost | None:
    """A live row of ``store`` other than ``exclude`` in effect on any day from
    ``start`` to ``end`` (open when ``None``), or ``None``."""
    rows = live_costs(store).exclude(pk=exclude.pk).exclude(effective_to__lt=start)
    if end is not None:
        rows = rows.filter(effective_from__lte=end)
    return rows.order_by("effective_from").first()


def correct_cost(
    store: Store,
    row: StoreCost,
    values: CostValues,
    effective_from: date,
    effective_to: date | None,
    note: str,
    actor: User | None,
) -> CostOutcome:
    """Correct one live cost row of ``store`` (admin, D10): a new row with ``values``
    and the dates, and the old row superseded by it with ``note``, in one transaction
    under the store's row lock. Refused, with nothing written, when the note is empty,
    the values break a rule, the end precedes the start, the row is not a live row of
    the store, or the new dates overlap another live row."""
    _require_admin(actor)
    text = visible_text(note or "")
    values = CostValues(
        landlord_rent=values.landlord_rent,
        other_monthly_cost=values.other_monthly_cost,
        other_cost_description=visible_text(values.other_cost_description or ""),
        landlord_lease_end=values.landlord_lease_end,
        notes=visible_text(values.notes or ""),
    )
    refused = _note_refusal(text, CHANGE_NOTE_REQUIRED) or values_refusal(values)
    if refused is not None:
        return refused
    if effective_to is not None and effective_to < effective_from:
        return CostOutcome(False, ACTION_REFUSED, END_BEFORE_START, "effective_to")
    try:
        with transaction.atomic():
            locked = Store.objects.select_for_update().get(pk=store.pk)
            old = StoreCost.objects.filter(store=locked, pk=row.pk).first()
            if old is None:
                return CostOutcome(False, ACTION_REFUSED, ROW_NOT_FOUND)
            if old.superseded_at is not None:
                return CostOutcome(False, ACTION_REFUSED, ALREADY_SUPERSEDED)
            if (
                values.same_as(old)
                and old.effective_from == effective_from
                and old.effective_to == effective_to
            ):
                return CostOutcome(True, ACTION_UNCHANGED, UNCHANGED)
            other = _overlapping(locked, effective_from, effective_to, exclude=old)
            if other is not None:
                return CostOutcome(
                    False,
                    ACTION_REFUSED,
                    OVERLAPS.format(
                        start=other.effective_from.isoformat(),
                        end=other.effective_to.isoformat() if other.effective_to else "open",
                    ),
                    "effective_from",
                )
            before = _state(old)
            new = StoreCost(
                store=locked,
                effective_from=effective_from,
                effective_to=effective_to,
                change_note=text,
            )
            _assign(new, values, StoreCostSource.STAFF)
            # The old row leaves the live set first, so the new one may take its start
            # date or be the open row; it names the new row, inserted next in the same
            # transaction (the foreign key is checked at commit).
            old.superseded_at = timezone.now()
            old.superseded_by_id = new.pk
            old.change_note = text
            old.save(update_fields=["superseded_at", "superseded_by", "change_note", "updated_at"])
            new.save(force_insert=True)
            _audit(actor, ACTION_AUDIT_SUPERSEDED, old, before, _state(old))
            _audit(actor, ACTION_AUDIT_CORRECTION, new, None, _state(new))
            return CostOutcome(True, ACTION_CREATED, CORRECTED, cost=new)
    except ValidationError:
        return CostOutcome(False, ACTION_REFUSED, PAYMENT_DATA_MESSAGE)
    except IntegrityError:
        # Not reachable under the store lock; kept so a constraint never becomes a 500.
        return CostOutcome(
            False, ACTION_REFUSED, "The cost row could not be corrected. Nothing was saved."
        )


def close_cost(store: Store, effective_to: date, note: str, actor: User | None) -> CostOutcome:
    """End the store's open cost row on ``effective_to`` (admin, D10), with ``note``.
    Later days have no cost row: the cost is unknown there, never zero."""
    _require_admin(actor)
    text = visible_text(note or "")
    refused = _note_refusal(text, CLOSE_NOTE_REQUIRED)
    if refused is not None:
        return refused
    try:
        with transaction.atomic():
            locked = Store.objects.select_for_update().get(pk=store.pk)
            row = open_cost(locked)
            if row is None:
                return CostOutcome(False, ACTION_REFUSED, NO_OPEN_ROW)
            if effective_to < row.effective_from:
                return CostOutcome(
                    False,
                    ACTION_REFUSED,
                    CLOSE_BEFORE_START.format(start=row.effective_from.isoformat()),
                    "effective_to",
                )
            before = _state(row)
            row.effective_to = effective_to
            row.change_note = text
            row.save(update_fields=["effective_to", "change_note", "updated_at"])
            _audit(actor, ACTION_AUDIT_CLOSED_BY_ADMIN, row, before, _state(row))
            return CostOutcome(
                True, ACTION_UPDATED, CLOSED.format(end=effective_to.isoformat()), cost=row
            )
    except ValidationError:
        return CostOutcome(False, ACTION_REFUSED, PAYMENT_DATA_MESSAGE)


# --- Import: reading the file (D10) ---------------------------------------------------

SHEET_NAME = "Store costs"
#: The columns the importer reads; every one must be in the header.
REQUIRED_COLUMNS: tuple[str, ...] = (
    "store_no",
    "walmart_monthly_rent",
    "walmart_rent_effective_from",
    "walmart_lease_end",
    "other_monthly_cost",
    "other_cost_description",
    "notes",
)
#: Columns copied from the Rents tab for reference; present or missing, never read.
REFERENCE_COLUMNS: tuple[str, ...] = (
    "location",
    "state",
    "operator (ref only)",
    "operator rent (ref only)",
)
#: Non-empty data rows (after the header) a file may have.
MAX_ROWS = 10_000
#: Lines a file may have, empty ones included: reading stops there.
MAX_LINES = 10 * MAX_ROWS
#: Largest file accepted, and largest total of a workbook's unpacked parts.
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNPACKED_BYTES = 100 * 1024 * 1024
NOTES_MAX_LENGTH = 2000

FILE_TYPE = "Only .csv and .xlsx files can be imported. Nothing was imported."
FILE_UNREADABLE = "The file could not be read. Nothing was imported."
FILE_TOO_LARGE = "The file is larger than 20 MB. Nothing was imported."
NOT_CSV = (
    'The file is not a CSV in UTF-8 text: save it as "CSV UTF-8" or as .xlsx. Nothing was imported.'
)
NOT_WORKBOOK = "The file is not a valid .xlsx workbook. Nothing was imported."
NO_SHEET = f'The workbook has no sheet named "{SHEET_NAME}". Nothing was imported.'
NO_HEADER = "The file has no header row. Nothing was imported."
MISSING_COLUMNS = "The header has no column {names}. Nothing was imported."
REPEATED_COLUMNS = "The header has the column {names} more than once. Nothing was imported."
TOO_MANY_ROWS = f"The file has more than {MAX_ROWS:,} rows. Nothing was imported."


class FileRefused(Exception):  # noqa: N818 (reads as an outcome, like SetupRefused)
    """The whole file is refused; ``message`` is one value-free sentence."""

    def __init__(self, message: str) -> None:
        super().__init__(message)
        self.message = message


@dataclass(frozen=True, slots=True)
class ParsedRow:
    """One non-empty data row: its row number in the file (the header is row 1) and the
    text of each named column."""

    number: int
    cells: Mapping[str, str]
    #: Columns whose cell is a formula with no saved value (``.xlsx`` only).
    formulas: frozenset[str] = frozenset()


@dataclass(frozen=True, slots=True)
class ParsedFile:
    #: The header's column names, normalised (stripped, lower case), in file order.
    columns: tuple[str, ...]
    rows: tuple[ParsedRow, ...]


def parse_file(path: str | Path) -> ParsedFile:
    """Read a ``.csv`` or ``.xlsx`` store-costs file; :class:`FileRefused` when the file
    as a whole cannot be imported."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        raise FileRefused(FILE_TYPE)
    try:
        size = path.stat().st_size
    except OSError:
        raise FileRefused(FILE_UNREADABLE) from None
    if size > MAX_FILE_BYTES:
        raise FileRefused(FILE_TOO_LARGE)
    return _read_csv(path) if suffix == ".csv" else _read_xlsx(path)


#: One line of a file: its cells, and the positions of formula cells with no saved value.
Line = tuple[list[Any], frozenset[int]]


def _read_csv(path: Path) -> ParsedFile:
    try:
        text = path.read_bytes().decode("utf-8-sig")
    except OSError:
        raise FileRefused(FILE_UNREADABLE) from None
    except UnicodeDecodeError:
        raise FileRefused(NOT_CSV) from None
    if "\x00" in text:
        raise FileRefused(NOT_CSV)
    try:
        return _parse((list(line), frozenset()) for line in csv.reader(io.StringIO(text)))
    except csv.Error:
        raise FileRefused(NOT_CSV) from None


def _read_xlsx(path: Path) -> ParsedFile:
    """The "Store costs" sheet as cached values. Two read-only passes over the sheet,
    side by side: one for the saved values, one to tell a formula from an empty cell."""
    try:
        if not zipfile.is_zipfile(path):
            raise FileRefused(NOT_WORKBOOK)
        with zipfile.ZipFile(path) as archive:
            if sum(member.file_size for member in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise FileRefused(NOT_WORKBOOK)
    except (OSError, zipfile.BadZipFile):
        raise FileRefused(NOT_WORKBOOK) from None
    import openpyxl  # type: ignore[import-untyped]  # the application's only import of it

    with warnings.catch_warnings():
        # Unsupported styles or extensions in a workbook are not the import's concern.
        warnings.simplefilter("ignore")
        books = []
        try:
            for data_only in (True, False):
                books.append(
                    openpyxl.load_workbook(
                        path, read_only=True, data_only=data_only, keep_links=False
                    )
                )
            values_book, formula_book = books
            if SHEET_NAME not in values_book.sheetnames:
                raise FileRefused(NO_SHEET)
            lines = (
                (
                    [cell.value for cell in value_row],
                    frozenset(
                        position
                        for position, (value, formula) in enumerate(
                            zip(value_row, formula_row, strict=False)
                        )
                        if getattr(formula, "data_type", "") == "f" and value.value is None
                    ),
                )
                for value_row, formula_row in zip(
                    values_book[SHEET_NAME].iter_rows(),
                    formula_book[SHEET_NAME].iter_rows(),
                    strict=False,
                )
            )
            return _parse(lines)
        except FileRefused:
            raise
        except Exception:  # any reading failure is the same sentence; nothing leaks
            raise FileRefused(NOT_WORKBOOK) from None
        finally:
            for book in books:
                book.close()


def _parse(lines: Iterable[Line]) -> ParsedFile:
    """The header, then every non-empty row; refused past :data:`MAX_ROWS` non-empty
    rows or :data:`MAX_LINES` lines, without reading further."""
    header: tuple[str, ...] | None = None
    rows: list[ParsedRow] = []
    number = 0
    for raw, uncached in lines:
        number += 1
        texts = [cell_text(value) for value in raw]
        if header is None:
            header = _check_header(texts)
            continue
        if number > MAX_LINES:
            raise FileRefused(TOO_MANY_ROWS)
        if not any(text.strip() for text in texts) and not uncached:
            continue
        if len(rows) == MAX_ROWS:
            raise FileRefused(TOO_MANY_ROWS)
        cells = {name: (texts[i] if i < len(texts) else "") for i, name in enumerate(header)}
        formulas = frozenset(header[i] for i in uncached if i < len(header))
        rows.append(ParsedRow(number=number, cells=cells, formulas=formulas))
    if header is None:
        raise FileRefused(NO_HEADER)
    return ParsedFile(columns=header, rows=tuple(rows))


def _check_header(texts: list[str]) -> tuple[str, ...]:
    names = tuple(text.strip().lower() for text in texts)
    if not any(names):
        raise FileRefused(NO_HEADER)
    missing = [name for name in REQUIRED_COLUMNS if name not in names]
    if missing:
        raise FileRefused(MISSING_COLUMNS.format(names=", ".join(missing)))
    counts = Counter(names)
    repeated = [name for name in REQUIRED_COLUMNS if counts[name] > 1]
    if repeated:
        raise FileRefused(REPEATED_COLUMNS.format(names=", ".join(repeated)))
    return names


def cell_text(value: Any) -> str:
    """A cell as text: a CSV field as it is; a workbook number without a needless
    ``.0`` (``2500.0`` is ``2500``); a workbook date as ``YYYY-MM-DD`` (with the time
    when it has one, which no date column accepts)."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, numbers.Real):
        # A workbook number is a binary fraction: read through its shortest text, never
        # through its binary value, so 2500.1 stays 2500.1.
        exact = Decimal(repr(value))
        if exact.is_finite() and exact == exact.to_integral_value():
            return str(int(exact))
        return format(exact, "f") if exact.is_finite() else str(exact)
    return str(value)


# --- Import: reading one row ----------------------------------------------------------

_AMOUNT = re.compile(r"[0-9]+(?:\.[0-9]{1,2})?")
_THREE_DECIMALS = re.compile(r"[0-9]+\.[0-9]{3,}")
_STORE_NO = re.compile(r"[0-9]{1,10}")
_US_DATE = re.compile(r"([0-9]{1,2})/([0-9]{1,2})/([0-9]{4})")
_ISO_DATE = re.compile(r"([0-9]{4})-([0-9]{2})-([0-9]{2})")
INT_MAX = 2_147_483_647

ROW_PAYMENT_DATA = "a cell may hold a card or account number"
ROW_FORMULA = "{column} holds a formula with no saved value; type the value instead"
ROW_STORE_NO = "store_no is not a store number"
ROW_AMOUNT = (
    "{column} is not a plain number (digits with an optional decimal point; "
    "no $, commas or minus sign)"
)
ROW_DECIMALS = "{column} has more than two decimals"
ROW_DATE = "{column} is not a date (MM/DD/YYYY or YYYY-MM-DD)"
ROW_TOO_LONG = "{column} is longer than {limit} characters"
ROW_UNKNOWN_STORE = "no store has number {number}"
ROW_MANY_STORES = "more than one store has number {number}"
ROW_DUPLICATE = "store {number} appears more than once in the file"
ROW_NO_OPEN_DATE = (
    "walmart_rent_effective_from is empty and store {number} has no open date; "
    "fill walmart_rent_effective_from"
)


@dataclass(frozen=True, slots=True)
class RowReading:
    """One row read: ``not_filled``, or ``refusal`` (a value-free reason), or values."""

    not_filled: bool = False
    refusal: str = ""
    store_number: int | None = None
    values: CostValues | None = None
    effective_from: date | None = None


def read_row(row: ParsedRow) -> RowReading:
    """Check one row on its own (no database read): the card-number guard on every
    cell, formulas, then each column the importer uses."""
    for text in row.cells.values():
        try:
            validate_no_payment_data(text)
        except ValidationError:
            return RowReading(refusal=ROW_PAYMENT_DATA)
    used = [name for name in REQUIRED_COLUMNS if name in row.formulas]
    if used:
        return RowReading(refusal=ROW_FORMULA.format(column=used[0]))
    cells = {name: row.cells.get(name, "").strip() for name in REQUIRED_COLUMNS}
    if not cells["walmart_monthly_rent"]:
        return RowReading(not_filled=True)
    store_number = parse_store_number(cells["store_no"])
    if store_number is None:
        return RowReading(refusal=ROW_STORE_NO)
    refusal = ""
    rent = other = None
    rent, refusal = _amount(cells, "walmart_monthly_rent")
    if not refusal and cells["other_monthly_cost"]:
        other, refusal = _amount(cells, "other_monthly_cost")
    effective_from = lease_end = None
    if not refusal and cells["walmart_rent_effective_from"]:
        effective_from = parse_date(cells["walmart_rent_effective_from"])
        if effective_from is None:
            refusal = ROW_DATE.format(column="walmart_rent_effective_from")
    if not refusal and cells["walmart_lease_end"]:
        lease_end = parse_date(cells["walmart_lease_end"])
        if lease_end is None:
            refusal = ROW_DATE.format(column="walmart_lease_end")
    description = visible_text(cells["other_cost_description"])
    notes = visible_text(cells["notes"])
    if not refusal and len(description) > NAME_MAX_LENGTH:
        refusal = ROW_TOO_LONG.format(column="other_cost_description", limit=NAME_MAX_LENGTH)
    if not refusal and len(notes) > NOTES_MAX_LENGTH:
        refusal = ROW_TOO_LONG.format(column="notes", limit=NOTES_MAX_LENGTH)
    if refusal or rent is None:
        return RowReading(refusal=refusal, store_number=store_number)
    return RowReading(
        store_number=store_number,
        effective_from=effective_from,
        values=CostValues(
            landlord_rent=rent,
            other_monthly_cost=other,
            other_cost_description=description,
            landlord_lease_end=lease_end,
            notes=notes,
        ),
    )


def parse_store_number(text: str) -> int | None:
    """``store_no`` as the store number: digits only, spaces around and leading zeros
    ignored (``" 01796 "`` is 1796)."""
    text = text.strip()
    if not _STORE_NO.fullmatch(text):
        return None
    number = int(text)
    return number if 0 < number <= INT_MAX else None


def parse_amount(text: str) -> Decimal | None:
    """A plain amount: digits with an optional decimal point and one or two decimals."""
    text = text.strip()
    if not _AMOUNT.fullmatch(text):
        return None
    value = Decimal(text).quantize(CENT)
    return value if value <= MONEY_MAX else None


def parse_date(text: str) -> date | None:
    """``MM/DD/YYYY`` (one-digit month and day allowed) or ``YYYY-MM-DD``; ``None`` for
    anything else, or for a day that does not exist (``02/30/2026``)."""
    text = text.strip()
    us = _US_DATE.fullmatch(text)
    iso = _ISO_DATE.fullmatch(text)
    try:
        if us:
            return date(int(us.group(3)), int(us.group(1)), int(us.group(2)))
        if iso:
            return date(int(iso.group(1)), int(iso.group(2)), int(iso.group(3)))
    except ValueError:
        return None
    return None


def _amount(cells: Mapping[str, str], column: str) -> tuple[Decimal | None, str]:
    text = cells[column]
    if _THREE_DECIMALS.fullmatch(text):
        return None, ROW_DECIMALS.format(column=column)
    value = parse_amount(text)
    if value is None:
        return None, ROW_AMOUNT.format(column=column)
    return value, ""


# --- Import: applying the rows --------------------------------------------------------


@dataclass
class ImportReport:
    read: int = 0
    not_filled: int = 0
    unchanged: int = 0
    created: int = 0
    updated: int = 0
    #: ``(row number, reason)`` per refused row, in file order.
    refusals: list[tuple[int, str]] = field(default_factory=list)

    @property
    def refused(self) -> int:
        return len(self.refusals)

    def counts(self) -> str:
        return (
            f"read {self.read}, not filled {self.not_filled}, unchanged {self.unchanged}, "
            f"created {self.created}, updated {self.updated}, refused {self.refused}"
        )


def import_rows(parsed: ParsedFile, *, dry_run: bool, actor: User | None) -> ImportReport:
    """Apply each row of ``parsed`` (D10), or with ``dry_run`` only say what would be
    done. Each row is applied in its own transaction; a refused row writes nothing."""
    report = ImportReport()
    readings = [(row, read_row(row)) for row in parsed.rows]
    filled = Counter(
        reading.store_number
        for _, reading in readings
        if not reading.not_filled and reading.store_number is not None
    )
    for row, reading in readings:
        report.read += 1
        if reading.not_filled:
            report.not_filled += 1
            continue
        if reading.store_number is not None and filled[reading.store_number] > 1:
            report.refusals.append((row.number, ROW_DUPLICATE.format(number=reading.store_number)))
            continue
        if reading.refusal or reading.values is None or reading.store_number is None:
            report.refusals.append((row.number, reading.refusal or ROW_STORE_NO))
            continue
        outcome = _apply(parsed, row, reading, dry_run=dry_run, actor=actor)
        if not outcome.ok:
            report.refusals.append((row.number, _row_reason(outcome)))
        elif outcome.action == ACTION_UNCHANGED:
            report.unchanged += 1
        elif outcome.action == ACTION_UPDATED:
            report.updated += 1
        else:
            report.created += 1
    return report


def _apply(
    parsed: ParsedFile,
    row: ParsedRow,
    reading: RowReading,
    *,
    dry_run: bool,
    actor: User | None,
) -> CostOutcome:
    number = reading.store_number
    values = reading.values
    if number is None or values is None:
        return CostOutcome(False, ACTION_REFUSED, ROW_STORE_NO)
    stores = list(Store.objects.filter(store_number=number)[:2])
    if not stores:
        return CostOutcome(False, ACTION_REFUSED, ROW_UNKNOWN_STORE.format(number=number))
    if len(stores) > 1:
        return CostOutcome(False, ACTION_REFUSED, ROW_MANY_STORES.format(number=number))
    store = stores[0]
    effective_from = reading.effective_from or store.open_date
    if effective_from is None:
        return CostOutcome(False, ACTION_REFUSED, ROW_NO_OPEN_DATE.format(number=number))
    if dry_run:
        return decide(latest_cost(store), values, effective_from)
    with transaction.atomic():
        outcome = set_cost(
            store,
            values.landlord_rent,
            effective_from,
            other_monthly_cost=values.other_monthly_cost,
            other_cost_description=values.other_cost_description,
            landlord_lease_end=values.landlord_lease_end,
            notes=values.notes,
            source=StoreCostSource.IMPORT,
            actor=actor,
        )
        if outcome.ok and outcome.cost is not None:
            rule = outcome.action
            if outcome.closed is not None:
                rule += f"; closed the row from {outcome.closed.effective_from.isoformat()}"
            if reading.effective_from is None:
                rule += "; effective from the store's open date"
            ImportAudit.objects.create(
                sheet_row=row.number,
                sheet_col=column_letter(parsed.columns.index("walmart_monthly_rent")),
                raw_value="; ".join(
                    f"{name}={row.cells.get(name, '')}" for name in REQUIRED_COLUMNS
                ),
                target_table=TARGET_TABLE,
                target_id=str(outcome.cost.pk),
                rule_applied=rule,
            )
    return outcome


def _row_reason(outcome: CostOutcome) -> str:
    """A service sentence as a row reason: no trailing "Nothing was saved."."""
    text = outcome.message.replace(" Nothing was saved.", "").rstrip(".")
    return text[:1].lower() + text[1:] if text else ACTION_REFUSED


def column_letter(index: int) -> str:
    """The spreadsheet letter of a zero-based column index (0 is A, 26 is AA)."""
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters
