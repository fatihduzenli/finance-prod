"""Store costs and their importer (docs/qa/PLAN-010b-setup-fees-store-costs-pnl.md Batch B:
D8, D10, D12; §6 test plan).

The date-effective service (no overlap, unknown never zero, the margin), the importer
for the owner's template (file checks, row rules, idempotency, dry run, audit), the
``import_store_costs`` command, the store page and its "Set landlord rent" form, and the
role gates. No call reaches the payment provider. ``openpyxl`` is used here only to
build workbooks for the reader under test.
"""

from __future__ import annotations

import ast
import csv
import io
import time
import tomllib
import uuid
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]
import pytest
from django.core.management import CommandError, call_command
from django.db import connection
from django.test.utils import CaptureQueriesContext

from accounts import registry
from accounts.roles import Role
from ledger.models import (
    AuditLog,
    ImportAudit,
    Lease,
    Store,
    StoreCost,
    StoreCostSource,
    StoreStatus,
)
from ledger.services import leases
from ledger.services import store_costs as service
from tests.conftest import REPO_ROOT
from tests.test_lease_services import _race, make_lease, make_operator
from tests.test_operator_views import client_for

pytestmark = pytest.mark.django_db

requires_postgresql = pytest.mark.skipif(
    connection.vendor != "postgresql",
    reason="row locks and concurrent transactions need PostgreSQL",
)

D = Decimal
FIXTURES = Path(__file__).parent / "fixtures" / "store_costs"
EXAMPLES = REPO_ROOT / "docs" / "templates" / "examples"
TEMPLATES = REPO_ROOT / "docs" / "templates"
HEADER = [
    "store_no",
    "location",
    "state",
    "operator (ref only)",
    "operator rent (ref only)",
    "walmart_monthly_rent",
    "walmart_rent_effective_from",
    "walmart_lease_end",
    "other_monthly_cost",
    "other_cost_description",
    "notes",
]
CARD = "4111 1111 1111 1111"
JAN = date(2026, 1, 1)


# --- Builders ----------------------------------------------------------------------------


def store(number: int | None = None, open_date: date | None = date(2025, 6, 1)) -> Store:
    return Store.objects.create(
        external_label=f"Store {number or uuid.uuid4().hex[:6]}",
        store_number=number,
        open_date=open_date,
        status=StoreStatus.ACTIVE,
    )


def set_cost(target: Store, rent: str, starts: date = JAN, **fields: Any) -> service.CostOutcome:
    return service.set_cost(target, D(rent), starts, **fields)


def rows_of(target: Store) -> list[tuple[Decimal, date, date | None]]:
    return [
        (row.landlord_rent, row.effective_from, row.effective_to)
        for row in StoreCost.objects.filter(store=target).order_by("effective_from")
    ]


def write_csv(path: Path, rows: list[list[str]], header: list[str] | None = None) -> Path:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header or HEADER)
    writer.writerows(rows)
    path.write_text(buffer.getvalue(), encoding="utf-8")
    return path


def row(store_no: str, rent: str, starts: str = "01/01/2026", **cells: str) -> list[str]:
    values = dict.fromkeys(HEADER, "")
    values.update(store_no=store_no, walmart_monthly_rent=rent, walmart_rent_effective_from=starts)
    values.update(cells)
    return [values[name] for name in HEADER]


def run_import(path: Path, *, dry_run: bool = False) -> service.ImportReport:
    return service.import_rows(service.parse_file(path), dry_run=dry_run, actor=None)


def refused_file(path: Path) -> str:
    with pytest.raises(service.FileRefused) as raised:
        service.parse_file(path)
    return raised.value.message


# --- The service: date-effective rows (D8) ---------------------------------------------


def test_a_first_cost_opens_one_row_and_is_audited() -> None:
    target = store()
    outcome = set_cost(target, "2500.00", other_monthly_cost=D("50.00"))
    assert (outcome.ok, outcome.action) == (True, service.ACTION_CREATED)
    assert rows_of(target) == [(D("2500.00"), JAN, None)]
    row = StoreCost.objects.get()
    assert (row.source, row.other_monthly_cost) == (StoreCostSource.STAFF, D("50.00"))
    entry = AuditLog.objects.get(action="store_cost.created")
    assert entry.entity_type == "store_cost" and entry.entity_id == str(row.pk)
    assert entry.after is not None and entry.after["landlord_rent"] == "2500.00"


def test_the_same_values_change_nothing_and_write_no_audit_row() -> None:
    target = store()
    set_cost(target, "2500.00")
    again = set_cost(target, "2500.00", starts=date(2026, 3, 1))
    assert (again.ok, again.action, again.message) == (True, "unchanged", service.UNCHANGED)
    assert len(rows_of(target)) == 1
    assert AuditLog.objects.count() == 1


def test_the_same_effective_date_updates_the_open_row_in_place() -> None:
    target = store()
    set_cost(target, "2500.00")
    outcome = set_cost(target, "2600.00", notes="Corrected from the lease.")
    assert outcome.action == service.ACTION_UPDATED
    assert rows_of(target) == [(D("2600.00"), JAN, None)]
    entry = AuditLog.objects.get(action="store_cost.updated")
    assert entry.before is not None and entry.before["landlord_rent"] == "2500.00"


def test_a_later_date_closes_the_open_row_the_day_before() -> None:
    target = store()
    set_cost(target, "2500.00")
    outcome = set_cost(target, "2700.00", starts=date(2026, 7, 1))
    assert outcome.action == service.ACTION_CREATED and outcome.closed is not None
    assert rows_of(target) == [
        (D("2500.00"), JAN, date(2026, 6, 30)),
        (D("2700.00"), date(2026, 7, 1), None),
    ]
    assert {"store_cost.closed", "store_cost.created"} <= set(
        AuditLog.objects.values_list("action", flat=True)
    )


def test_the_next_day_leaves_a_one_day_row() -> None:
    target = store()
    set_cost(target, "2500.00")
    set_cost(target, "2600.00", starts=JAN + timedelta(days=1))
    assert rows_of(target)[0] == (D("2500.00"), JAN, JAN)


def test_an_earlier_date_than_the_open_row_is_refused_and_writes_nothing() -> None:
    target = store()
    set_cost(target, "2500.00", starts=date(2026, 3, 1))
    outcome = set_cost(target, "2400.00", starts=JAN)
    assert (outcome.ok, outcome.field) == (False, "effective_from")
    assert outcome.message == service.EARLIER_THAN_OPEN.format(start="Mar 1, 2026")
    assert rows_of(target) == [(D("2500.00"), date(2026, 3, 1), None)]
    assert AuditLog.objects.count() == 1


def test_rows_never_overlap_and_leave_no_gap_whatever_the_order_of_changes() -> None:
    target = store()
    starts = [date(2026, 1, 1), date(2026, 1, 2), date(2026, 2, 28), date(2026, 3, 1)]
    for index, day in enumerate(starts):
        set_cost(target, f"{2500 + index}.00", starts=day)
        set_cost(target, f"{2600 + index}.00", starts=day)  # in place
    set_cost(target, "1.00", starts=date(2026, 1, 15))  # earlier than the open row: refused
    rows = rows_of(target)
    assert [start for _, start, _ in rows] == starts
    for (_, _, end), (_, next_start, _) in zip(rows, rows[1:], strict=False):
        assert end == next_start - timedelta(days=1)
    assert rows[-1][2] is None
    assert StoreCost.objects.filter(store=target, effective_to__isnull=True).count() == 1


@pytest.mark.parametrize(
    ("fields", "field", "message"),
    [
        ({"landlord_rent": D("-0.01")}, "landlord_rent", service.RENT_NEGATIVE),
        ({"landlord_rent": D("2500.001")}, "landlord_rent", service.NOT_CENTS),
        ({"other_monthly_cost": D("-1.00")}, "other_monthly_cost", service.OTHER_NEGATIVE),
        ({"other_cost_description": CARD}, "other_cost_description", None),
        ({"notes": f"Card {CARD}"}, "notes", None),
        ({"other_cost_description": "x" * 201}, "other_cost_description", None),
    ],
    ids=["negative", "three decimals", "negative other", "card", "card in notes", "too long"],
)
def test_bad_values_are_refused_before_anything_is_written(
    fields: dict[str, Any], field: str, message: str | None
) -> None:
    target = store()
    values: dict[str, Any] = {"landlord_rent": D("2500.00"), "effective_from": JAN, **fields}
    outcome = service.set_cost(target, **values)
    assert (outcome.ok, outcome.field) == (False, field)
    if message is not None:
        assert outcome.message == message
    assert not StoreCost.objects.exists() and not AuditLog.objects.exists()


def test_a_zero_landlord_rent_is_a_known_cost_not_an_unknown_one() -> None:
    target = store()
    set_cost(target, "0.00")
    assert service.current_cost(target, JAN) is not None


def test_current_cost_is_the_row_in_effect_that_day_or_none() -> None:
    target = store()
    set_cost(target, "2500.00")
    set_cost(target, "2700.00", starts=date(2026, 7, 1))
    assert service.current_cost(target, date(2025, 12, 31)) is None
    first = service.current_cost(target, date(2026, 6, 30))
    second = service.current_cost(target, date(2026, 7, 1))
    assert first is not None and first.landlord_rent == D("2500.00")
    assert second is not None and second.landlord_rent == D("2700.00")


@requires_postgresql
@pytest.mark.django_db(transaction=True)
def test_two_concurrent_sets_leave_one_open_row_and_no_overlap() -> None:
    target = store()
    set_cost(target, "2500.00")
    results = _race(
        lambda: set_cost(target, "2600.00", starts=date(2026, 5, 1)),
        lambda: set_cost(target, "2700.00", starts=date(2026, 5, 1)),
    )
    assert sorted(result.action for result in results) == ["created", "updated"]
    rows = rows_of(target)
    assert [(start, end) for _, start, end in rows] == [
        (JAN, date(2026, 4, 30)),
        (date(2026, 5, 1), None),
    ]


# --- The margin (D8) ---------------------------------------------------------------------


def leased(rent: str = "3500.00", **overrides: Any) -> Lease:
    tag = uuid.uuid4().hex[:8]
    values: dict[str, Any] = {
        "operator": make_operator(email=f"{tag}@example.com", customer=None),
        "started_on": JAN,
    }
    ended_on = overrides.pop("ended_on", None)
    values.update(overrides)
    lease = make_lease(D(rent), **values)
    if ended_on is not None:
        # Ended through the staff path, as the lease page does.
        lease = leases.end_lease(
            lease.pk, ended_on, "Operator moved out.", actor=None, today=ended_on
        )
    return lease


def test_margin_is_operator_rent_less_landlord_rent_less_other_cost() -> None:
    lease = leased("3500.00")
    set_cost(lease.store, "2500.00", other_monthly_cost=D("50.00"))
    assert service.operator_rent(lease.store, date(2026, 2, 1)) == D("3500.00")
    assert service.margin(lease.store, date(2026, 2, 1)) == D("950.00")


def test_an_empty_other_cost_means_none() -> None:
    lease = leased("3500.00")
    set_cost(lease.store, "2500.00")
    assert service.margin(lease.store, date(2026, 2, 1)) == D("1000.00")


def test_margin_is_unknown_without_a_cost_row_never_zero() -> None:
    lease = leased("3500.00")
    assert service.current_cost(lease.store, date(2026, 2, 1)) is None
    assert service.margin(lease.store, date(2026, 2, 1)) is None


def test_margin_is_unknown_without_an_active_lease() -> None:
    lease = leased("3500.00", ended_on=date(2026, 3, 31))
    set_cost(lease.store, "2500.00")
    assert service.margin(lease.store, date(2026, 3, 31)) == D("1000.00")
    assert service.margin(lease.store, date(2026, 4, 1)) is None
    alone = store()
    set_cost(alone, "2500.00")
    assert service.operator_rent(alone, date(2026, 2, 1)) is None
    assert service.margin(alone, date(2026, 2, 1)) is None


def test_margin_is_unknown_before_the_lease_starts_or_without_a_rent_row() -> None:
    lease = leased("3500.00", started_on=date(2026, 3, 1))
    set_cost(lease.store, "2500.00")
    assert service.margin(lease.store, date(2026, 2, 28)) is None
    lease.rent_history.all().delete()
    assert service.margin(lease.store, date(2026, 3, 15)) is None


def test_margin_follows_the_rent_and_the_cost_in_effect_that_day() -> None:
    lease = leased("3500.00")
    leases.update_lease(
        lease.pk,
        {name: getattr(lease, name) for name in leases.LEASE_FIELDS},
        rent=D("3700.00"),
        rent_effective_from=date(2026, 6, 1),
        actor=None,
    )
    set_cost(lease.store, "2500.00")
    set_cost(lease.store, "2550.00", starts=date(2026, 6, 15))
    assert service.margin(lease.store, date(2026, 5, 31)) == D("1000.00")
    assert service.margin(lease.store, date(2026, 6, 1)) == D("1200.00")
    assert service.margin(lease.store, date(2026, 6, 15)) == D("1150.00")


def test_on_a_hand_over_day_the_new_lease_counts() -> None:
    old = leased("3000.00", ended_on=date(2026, 5, 1))
    leased("3600.00", store=old.store, started_on=date(2026, 5, 1))
    set_cost(old.store, "2500.00")
    assert service.operator_rent(old.store, date(2026, 5, 1)) == D("3600.00")
    assert service.operator_rent(old.store, date(2026, 4, 30)) == D("3000.00")


# --- Reading a file (D10) ----------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "number"),
    [("1796", 1796), ("01796", 1796), (" 1796 ", 1796), ("1796 ", 1796), ("0", None)]
    + [("17 96", None), ("1796a", None), ("", None), ("-1796", None), ("99999999999", None)],
)
def test_store_no_is_digits_with_spaces_around_and_leading_zeros_ignored(
    text: str, number: int | None
) -> None:
    assert service.parse_store_number(text) == number


@pytest.mark.parametrize(
    ("text", "value"),
    [("2500", D("2500.00")), ("2500.5", D("2500.50")), ("2500.00", D("2500.00")), ("0", D(0))]
    + [("2,500", None), ("$2500", None), ("-1", None), ("2500.555", None), ("2500.", None)]
    + [(".5", None), ("1e3", None), ("２５００", None), ("99999999999", None)],
)
def test_amounts_are_plain_numbers_with_at_most_two_decimals(
    text: str, value: Decimal | None
) -> None:
    assert service.parse_amount(text) == value


@pytest.mark.parametrize(
    ("text", "value"),
    [("1/5/2026", date(2026, 1, 5)), ("01/05/2026", date(2026, 1, 5))]
    + [("2026-01-05", date(2026, 1, 5)), ("12/31/2030", date(2030, 12, 31))]
    + [("13/01/2026", None), ("02/30/2026", None), ("2026-02-30", None), ("1/5/26", None)]
    + [("2026/01/05", None), ("05.01.2026", None), ("", None)],
)
def test_dates_are_mm_dd_yyyy_or_iso_and_must_exist(text: str, value: date | None) -> None:
    assert service.parse_date(text) == value


def test_the_owners_template_parses_as_not_filled_in_both_formats() -> None:
    for name in ("Store_Costs_Template.csv", "Store_Costs_Template.xlsx"):
        parsed = service.parse_file(TEMPLATES / name)
        assert len(parsed.rows) == 517, name
        report = service.import_rows(parsed, dry_run=True, actor=None)
        assert (report.read, report.not_filled, report.refused) == (517, 517, 0), name


def test_both_example_files_parse_to_the_same_rows() -> None:
    readings = {}
    for suffix in ("csv", "xlsx"):
        parsed = service.parse_file(EXAMPLES / f"store_costs_two_rows.{suffix}")
        assert parsed.columns == tuple(HEADER)
        readings[suffix] = [service.read_row(row) for row in parsed.rows]
    assert readings["csv"] == readings["xlsx"]
    first, second = readings["csv"]
    assert (first.store_number, second.store_number) == (1796, 437)
    assert first.values is not None and second.values is not None
    assert (first.values.landlord_rent, first.effective_from) == (D("2500.00"), JAN)
    assert first.values.landlord_lease_end == date(2030, 12, 31)
    assert (second.values.landlord_rent, second.values.other_monthly_cost) == (
        D("3000.00"),
        D("50.00"),
    )
    assert second.values.other_cost_description == "dumpster"


def test_a_csv_with_a_byte_order_mark_is_read(tmp_path: Path) -> None:
    path = write_csv(tmp_path / "bom.csv", [row("1796", "2500")])
    path.write_bytes(b"\xef\xbb\xbf" + path.read_bytes())
    parsed = service.parse_file(path)
    assert parsed.columns[0] == "store_no"
    assert parsed.rows[0].cells["store_no"] == "1796"


def test_missing_reference_columns_are_accepted_and_extra_ones_ignored(tmp_path: Path) -> None:
    header = [*service.REQUIRED_COLUMNS, "my own column"]
    path = write_csv(tmp_path / "short.csv", [["1796", "2500", "", "", "", "", "", "x"]], header)
    parsed = service.parse_file(path)
    assert service.read_row(parsed.rows[0]).store_number == 1796


@pytest.mark.parametrize("missing", service.REQUIRED_COLUMNS)
def test_a_missing_required_column_refuses_the_whole_file(tmp_path: Path, missing: str) -> None:
    store(1796)
    header = [name for name in HEADER if name != missing]
    path = write_csv(tmp_path / "bad.csv", [["1796"] * len(header)], header)
    assert refused_file(path) == service.MISSING_COLUMNS.format(names=missing)
    assert not StoreCost.objects.exists()


def test_a_required_column_twice_refuses_the_whole_file(tmp_path: Path) -> None:
    path = write_csv(tmp_path / "twice.csv", [], [*HEADER, " Notes "])
    assert refused_file(path) == service.REPEATED_COLUMNS.format(names="notes")


def test_the_rents_tab_is_refused_by_its_header(tmp_path: Path) -> None:
    path = write_csv(tmp_path / "rents.csv", [["1796", "RASHAD", "3500"]], ["Store", "Op", "Rent"])
    assert refused_file(path).startswith("The header has no column store_no, ")


def test_ten_thousand_rows_are_read_and_one_more_refuses_the_file(tmp_path: Path) -> None:
    rows = [row(str(n), "") for n in range(1, service.MAX_ROWS + 1)]
    assert len(service.parse_file(write_csv(tmp_path / "max.csv", rows)).rows) == 10_000
    rows.append(row("10001", ""))
    assert refused_file(write_csv(tmp_path / "over.csv", rows)) == service.TOO_MANY_ROWS


def test_a_five_megabyte_file_of_a_hundred_thousand_rows_is_refused_quickly(
    tmp_path: Path,
) -> None:
    line = ",".join(row("1796", "2500.00", notes="x" * 20)) + "\n"
    path = tmp_path / "big.csv"
    path.write_text(",".join(HEADER) + "\n" + line * 100_000, encoding="utf-8")
    assert path.stat().st_size > 5_000_000
    started = time.monotonic()
    assert refused_file(path) == service.TOO_MANY_ROWS
    assert time.monotonic() - started < 10


def test_many_empty_lines_stop_the_reading(tmp_path: Path) -> None:
    path = tmp_path / "empty_lines.csv"
    path.write_text(",".join(HEADER) + "\n" + ",,,,,,,,,,\n" * (service.MAX_LINES + 1))
    assert refused_file(path) == service.TOO_MANY_ROWS


@pytest.mark.parametrize(
    ("name", "content", "message"),
    [
        ("binary.csv", b"\x89PNG\r\n\x1a\n\x00\x00\xff\xfe", service.NOT_CSV),
        ("nul.csv", b"store_no\x00,notes\n", service.NOT_CSV),
        ("latin1.csv", ",".join(HEADER).encode() + b"\n1796,Caf\xe9\n", service.NOT_CSV),
        ("empty.csv", b"", service.NO_HEADER),
        ("blank_header.csv", b",,,\n1796,2500\n", service.NO_HEADER),
        ("text.xlsx", b"store_no,walmart_monthly_rent\n1796,2500\n", service.NOT_WORKBOOK),
        ("random.xlsx", bytes(range(256)) * 4, service.NOT_WORKBOOK),
        ("notes.txt", b"hello", service.FILE_TYPE),
        ("no_extension", b"hello", service.FILE_TYPE),
        ("old.xls", b"hello", service.FILE_TYPE),
    ],
)
def test_a_file_that_is_not_a_csv_or_a_workbook_is_one_sentence(
    tmp_path: Path, name: str, content: bytes, message: str
) -> None:
    path = tmp_path / name
    path.write_bytes(content)
    assert refused_file(path) == message


def test_a_zip_that_is_not_a_workbook_is_one_sentence(tmp_path: Path) -> None:
    path = tmp_path / "archive.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")
    assert refused_file(path) == service.NOT_WORKBOOK


def test_a_workbook_that_unpacks_too_large_is_refused_before_it_is_opened(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * 2_000_000)
    assert path.stat().st_size < 100_000
    monkeypatch.setattr(service, "MAX_UNPACKED_BYTES", 1_000_000)
    monkeypatch.setattr(
        "openpyxl.load_workbook", lambda *args, **kwargs: pytest.fail("opened the bomb")
    )
    assert refused_file(path) == service.NOT_WORKBOOK


def test_a_file_above_the_size_limit_or_missing_is_one_sentence(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = write_csv(tmp_path / "big.csv", [row("1796", "2500")])
    monkeypatch.setattr(service, "MAX_FILE_BYTES", 10)
    assert refused_file(path) == service.FILE_TOO_LARGE
    assert refused_file(tmp_path / "missing.csv") == service.FILE_UNREADABLE


def test_a_workbook_without_the_store_costs_sheet_is_refused() -> None:
    assert refused_file(FIXTURES / "no_store_costs_sheet.xlsx") == service.NO_SHEET


def test_the_sheet_is_found_by_name_and_cells_read_as_typed() -> None:
    parsed = service.parse_file(FIXTURES / "typed_cells.xlsx")
    readings = {row.number: service.read_row(row) for row in parsed.rows}
    assert openpyxl.load_workbook(FIXTURES / "typed_cells.xlsx").sheetnames[0] == "How to fill"
    numeric = readings[2]  # numbers and date cells
    assert numeric.store_number == 1796 and numeric.values is not None
    assert (numeric.values.landlord_rent, numeric.effective_from) == (D("2500.00"), JAN)
    assert numeric.values.landlord_lease_end == date(2030, 12, 31)
    text = readings[3]  # text cells, " 0437 ", "1/5/2026", a numeric 50.5
    assert text.store_number == 437 and text.values is not None
    assert (text.values.landlord_rent, text.effective_from) == (D("3000.00"), date(2026, 1, 5))
    assert text.values.other_monthly_cost == D("50.50")
    saved_formula = readings[4]  # =2000+500 with its saved value
    assert saved_formula.values is not None and saved_formula.values.landlord_rent == D("2500.00")
    assert readings[5].refusal == service.ROW_FORMULA.format(column="walmart_monthly_rent")
    assert readings[6].values is not None and readings[6].values.landlord_rent == D("2500.50")
    assert readings[7].refusal == service.ROW_DECIMALS.format(column="walmart_monthly_rent")


def test_a_formula_is_never_evaluated(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Store costs"
    sheet.append(HEADER)
    sheet.append(["1796", None, None, None, None, "=1+1", "01/01/2026", None, None, None, None])
    path = tmp_path / "formula.xlsx"
    workbook.save(path)
    [parsed] = service.parse_file(path).rows
    assert parsed.cells["walmart_monthly_rent"] == "" and parsed.formulas == {
        "walmart_monthly_rent"
    }
    assert service.read_row(parsed).refusal.startswith("walmart_monthly_rent holds a formula")


@pytest.mark.parametrize(
    ("value", "text"),
    [(2500, "2500"), (2500.0, "2500"), (2500.5, "2500.5"), (0.1 + 0.2, "0.30000000000000004")]
    + [(True, "TRUE"), (None, ""), (date(2026, 1, 5), "2026-01-05")]
    + [(D("12.50"), "12.50")],
    ids=repr,
)
def test_workbook_cells_become_text_without_binary_rounding(value: object, text: str) -> None:
    assert service.cell_text(value) == text


# --- Applying the rows (D10) -------------------------------------------------------------


def example_stores() -> tuple[Store, Store]:
    return store(1796), store(437)


@pytest.mark.parametrize("suffix", ["csv", "xlsx"])
def test_the_example_creates_two_rows_and_a_second_run_changes_nothing(suffix: str) -> None:
    amherst, sealy = example_stores()
    path = EXAMPLES / f"store_costs_two_rows.{suffix}"
    first = run_import(path)
    assert (first.read, first.created, first.refused) == (2, 2, 0)
    assert rows_of(amherst) == [(D("2500.00"), JAN, None)]
    assert rows_of(sealy) == [(D("3000.00"), JAN, None)]
    cost = service.current_cost(sealy, JAN)
    assert cost is not None and cost.source == StoreCostSource.IMPORT
    assert (cost.other_monthly_cost, cost.other_cost_description) == (D("50.00"), "dumpster")
    audits = AuditLog.objects.count(), ImportAudit.objects.count()
    second = run_import(path)
    assert (second.read, second.unchanged, second.created, second.updated) == (2, 2, 0, 0)
    assert (AuditLog.objects.count(), ImportAudit.objects.count()) == audits


def test_the_csv_then_the_workbook_changes_nothing() -> None:
    example_stores()
    run_import(EXAMPLES / "store_costs_two_rows.csv")
    assert run_import(EXAMPLES / "store_costs_two_rows.xlsx").unchanged == 2


def test_each_applied_row_writes_one_import_audit_row() -> None:
    amherst, _ = example_stores()
    run_import(EXAMPLES / "store_costs_two_rows.csv")
    audit = ImportAudit.objects.get(sheet_row=2)
    cost = service.open_cost(amherst)
    assert cost is not None
    assert (audit.target_table, audit.target_id, audit.sheet_col) == (
        "store_costs",
        str(cost.pk),
        "F",
    )
    assert audit.rule_applied == "created"
    assert "walmart_monthly_rent=2500.00" in audit.raw_value
    assert ImportAudit.objects.count() == 2


def test_a_later_date_closes_the_row_the_same_date_updates_an_earlier_one_is_refused(
    tmp_path: Path,
) -> None:
    target = store(1796)
    run_import(write_csv(tmp_path / "a.csv", [row("1796", "2500.00", "01/01/2026")]))
    later = run_import(write_csv(tmp_path / "b.csv", [row("1796", "2600", "2026-07-01")]))
    assert later.created == 1
    assert ImportAudit.objects.order_by("-created_at").first().rule_applied == (  # type: ignore[union-attr]
        "created; closed the row from 2026-01-01"
    )
    same = run_import(write_csv(tmp_path / "c.csv", [row("1796", "2650", "07/01/2026")]))
    assert same.updated == 1
    earlier = run_import(write_csv(tmp_path / "d.csv", [row("1796", "2400", "06/01/2026")]))
    assert earlier.refusals == [
        (2, "the effective date is before Jul 1, 2026, when the current landlord rent started")
    ]
    assert rows_of(target) == [
        (D("2500.00"), JAN, date(2026, 6, 30)),
        (D("2650.00"), date(2026, 7, 1), None),
    ]


def test_an_empty_effective_date_uses_the_store_open_date(tmp_path: Path) -> None:
    target = store(1796, open_date=date(2024, 3, 15))
    report = run_import(write_csv(tmp_path / "a.csv", [row("1796", "2500", "")]))
    assert report.created == 1
    assert rows_of(target) == [(D("2500.00"), date(2024, 3, 15), None)]
    assert ImportAudit.objects.get().rule_applied == "created; effective from the store's open date"


def refused_rows_world() -> dict[int, Store]:
    stores = {number: store(number) for number in (1796, 437, *range(2101, 2114))}
    Store.objects.filter(pk=stores[2107].pk).update(open_date=None, needs_open_date=True)
    set_cost(stores[2108], "2500.00", starts=date(2026, 3, 1))
    return stores


EXPECTED_REFUSALS = [
    (3, "no store has number 999999"),
    (4, "store 437 appears more than once in the file"),
    (5, "store 437 appears more than once in the file"),
    (6, service.ROW_AMOUNT.format(column="walmart_monthly_rent")),
    (7, service.ROW_AMOUNT.format(column="walmart_monthly_rent")),
    (8, service.ROW_AMOUNT.format(column="walmart_monthly_rent")),
    (9, service.ROW_DECIMALS.format(column="walmart_monthly_rent")),
    (10, service.ROW_DATE.format(column="walmart_rent_effective_from")),
    (11, service.ROW_DATE.format(column="walmart_rent_effective_from")),
    (12, service.ROW_NO_OPEN_DATE.format(number=2107)),
    (13, "the effective date is before Mar 1, 2026, when the current landlord rent started"),
    (14, service.ROW_PAYMENT_DATA),
    (15, service.ROW_STORE_NO),
    (16, service.ROW_AMOUNT.format(column="other_monthly_cost")),
    (17, service.ROW_DATE.format(column="walmart_lease_end")),
    (18, service.ROW_PAYMENT_DATA),
]


def test_every_refusal_rule_refuses_its_row_and_the_rest_is_applied() -> None:
    stores = refused_rows_world()
    report = run_import(FIXTURES / "refused_rows.csv")
    assert report.refusals == EXPECTED_REFUSALS
    assert (report.read, report.not_filled, report.created, report.updated) == (18, 1, 1, 0)
    assert rows_of(stores[1796]) == [(D("2500.00"), JAN, None)]
    # Refused rows wrote nothing: one import_audit row (the applied one), and only the
    # two stores with costs have rows.
    assert ImportAudit.objects.count() == 1
    assert set(StoreCost.objects.values_list("store__store_number", flat=True)) == {1796, 2108}
    assert rows_of(stores[2108]) == [(D("2500.00"), date(2026, 3, 1), None)]


def test_no_reason_repeats_a_cell_value() -> None:
    refused_rows_world()
    report = run_import(FIXTURES / "refused_rows.csv")
    reasons = " ".join(reason for _, reason in report.refusals)
    for value in ("2,500", "$2500", "-1", "2500.555", "13/01/2026", "02/30/2026", "abc"):
        assert value not in reasons, value
    # The whole number and its spaced form, never a fragment.
    assert CARD not in reasons and CARD.replace(" ", "") not in reasons


def test_a_refused_row_is_written_in_its_own_transaction(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    store(1796)
    store(437)
    original = ImportAudit.objects.create

    def fail_for_437(**fields: Any) -> ImportAudit:
        if fields["sheet_row"] == 3:
            raise RuntimeError("audit failure")
        return original(**fields)

    monkeypatch.setattr(ImportAudit.objects, "create", fail_for_437)
    path = write_csv(tmp_path / "a.csv", [row("1796", "2500"), row("437", "3000")])
    with pytest.raises(RuntimeError):
        run_import(path)
    # Row 2 committed on its own; row 3's cost row was rolled back with its audit row.
    assert set(StoreCost.objects.values_list("store__store_number", flat=True)) == {1796}


def test_two_stores_with_the_same_number_refuse_the_row(tmp_path: Path) -> None:
    store(1796)
    store(1796)
    report = run_import(write_csv(tmp_path / "a.csv", [row("1796", "2500")]))
    assert report.refusals == [(2, "more than one store has number 1796")]


def test_the_dry_run_reports_the_same_counts_and_writes_nothing() -> None:
    stores = refused_rows_world()
    set_cost(stores[1796], "2400.00", starts=date(2025, 12, 1))
    before = (StoreCost.objects.count(), AuditLog.objects.count(), ImportAudit.objects.count())
    with CaptureQueriesContext(connection) as queries:
        dry = run_import(FIXTURES / "refused_rows.csv", dry_run=True)
    writes = [
        query["sql"]
        for query in queries.captured_queries
        if query["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))
    ]
    assert writes == []
    assert (StoreCost.objects.count(), AuditLog.objects.count(), ImportAudit.objects.count()) == (
        before
    )
    real = run_import(FIXTURES / "refused_rows.csv")
    assert (dry.read, dry.not_filled, dry.unchanged, dry.created, dry.updated) == (
        real.read,
        real.not_filled,
        real.unchanged,
        real.created,
        real.updated,
    )
    assert dry.refusals == real.refusals and real.created == 1


# --- The command (D10) -------------------------------------------------------------------


def command(*args: str) -> str:
    out = io.StringIO()
    call_command("import_store_costs", *args, stdout=out)
    return out.getvalue()


def test_the_command_prints_the_counts_and_a_second_run_is_unchanged() -> None:
    example_stores()
    path = str(EXAMPLES / "store_costs_two_rows.xlsx")
    assert command(path).strip() == (
        "read 2, not filled 0, unchanged 0, created 2, updated 0, refused 0"
    )
    assert "unchanged 2, created 0" in command(path)


def test_the_command_lists_each_refusal_and_exits_1() -> None:
    refused_rows_world()
    out = io.StringIO()
    with pytest.raises(CommandError) as raised:
        call_command("import_store_costs", str(FIXTURES / "refused_rows.csv"), stdout=out)
    assert raised.value.returncode == 1
    lines = out.getvalue().splitlines()
    assert lines[0] == "read 18, not filled 1, unchanged 0, created 1, updated 0, refused 16"
    assert lines[1:] == [f"row {number}: {reason}" for number, reason in EXPECTED_REFUSALS]


def test_the_command_dry_run_says_so_and_writes_nothing() -> None:
    example_stores()
    out = command(str(EXAMPLES / "store_costs_two_rows.csv"), "--dry-run")
    assert out.startswith("dry run, nothing written: read 2, not filled 0, unchanged 0, created 2")
    assert not StoreCost.objects.exists() and not ImportAudit.objects.exists()


def test_the_command_refuses_a_whole_file_in_one_sentence(tmp_path: Path) -> None:
    path = tmp_path / "x.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(CommandError) as raised:
        command(str(path))
    assert (str(raised.value), raised.value.returncode) == (service.NOT_WORKBOOK, 1)


# --- The dependency (openpyxl, approved for this batch only) ------------------------------


def test_openpyxl_is_pinned_and_imported_by_the_importer_only() -> None:
    manifest = tomllib.loads((REPO_ROOT / "pyproject.toml").read_text())
    assert "openpyxl==3.1.5" in manifest["project"]["dependencies"]
    assert 'name = "openpyxl"\nversion = "3.1.5"' in (REPO_ROOT / "uv.lock").read_text()
    importers = set()
    for path in (REPO_ROOT / "backend").rglob("*.py"):
        relative = path.relative_to(REPO_ROOT).as_posix()
        if relative.startswith("backend/tests/"):
            continue
        for node in ast.walk(ast.parse(path.read_text())):
            names = []
            if isinstance(node, ast.Import):
                names = [alias.name for alias in node.names]
            elif isinstance(node, ast.ImportFrom) and node.module:
                names = [node.module]
            if any(name.split(".")[0] == "openpyxl" for name in names):
                importers.add(relative)
    assert importers == {"backend/ledger/services/store_costs.py"}


# --- The store page and "Set landlord rent" (D8, D12) --------------------------------------


def page(client: Any, target: Store) -> str:
    return str(client.get(f"/stores/{target.pk}/").content.decode())


def cost_form(**overrides: str) -> dict[str, str]:
    values = {
        "landlord_rent": "2500.00",
        "effective_from": "2026-01-01",
        "other_monthly_cost": "50.00",
        "other_cost_description": "dumpster",
        "landlord_lease_end": "2030-12-31",
        "notes": "",
    }
    return {**values, **overrides}


def test_the_store_page_shows_unknown_until_a_cost_is_set_then_the_cost() -> None:
    lease = leased("3500.00")
    client, user = client_for(Role.BILLING)
    shown = page(client, lease.store)
    assert '<dd id="landlord-rent">unknown</dd>' in shown
    # PLAN-012e-2 D8: the Store cost block shows no margin and no operator rent (the P&L
    # has the margin).
    assert 'id="store-margin"' not in shown and 'id="operator-rent"' not in shown
    response = client.post(f"/stores/{lease.store_id}/cost/", cost_form())
    assert response.status_code == 302 and response["Location"] == f"/stores/{lease.store_id}/"
    shown = page(client, lease.store)
    assert '<dd id="landlord-rent">$2,500.00 a month from Jan 1, 2026</dd>' in shown
    assert '<dd id="other-cost">$50.00 (dumpster)</dd>' in shown
    assert '<dd id="landlord-lease-end">Dec 31, 2030</dd>' in shown
    assert "Margin" not in shown and "$950.00" not in shown
    assert 'id="cost-history"' in shown
    entry = AuditLog.objects.get(action="store_cost.created")
    assert entry.actor == user


def test_the_margin_is_unknown_without_a_lease_and_never_on_the_page() -> None:
    target = store()
    set_cost(target, "2500.00")
    assert service.margin(target, date(2026, 6, 1)) is None
    client, _ = client_for(Role.VIEWER)
    shown = page(client, target)
    assert 'id="operator-rent"' not in shown and 'id="store-margin"' not in shown


def test_an_earlier_date_on_the_form_is_held_as_an_error_and_saves_nothing() -> None:
    target = store()
    set_cost(target, "2500.00", starts=date(2026, 3, 1))
    client, _ = client_for(Role.BILLING)
    path = f"/stores/{target.pk}/cost/"
    response = client.post(path, cost_form(effective_from="2026-01-01"))
    assert response["Location"] == path
    assert "The effective date is before Mar 1, 2026" in client.get(path).content.decode()
    assert StoreCost.objects.count() == 1


def test_a_card_number_on_the_form_is_refused_and_never_held() -> None:
    target = store()
    client, _ = client_for(Role.BILLING)
    path = f"/stores/{target.pk}/cost/"
    client.post(path, cost_form(other_cost_description=CARD, notes=f"pay {CARD}"))
    shown = client.get(path).content.decode()
    assert "may not contain card or account numbers" in shown
    # The whole number and its spaced form, never a fragment.
    assert CARD not in shown and CARD.replace(" ", "") not in shown
    assert not StoreCost.objects.exists()


def test_the_form_starts_from_the_open_row() -> None:
    target = store()
    set_cost(target, "2500.00", other_monthly_cost=D("75.00"), other_cost_description="parking")
    client, _ = client_for(Role.BILLING)
    shown = client.get(f"/stores/{target.pk}/cost/").content.decode()
    assert 'value="2500.00"' in shown and 'value="parking"' in shown
    assert "Current row: $2,500.00 a month from Jan 1, 2026." in shown


def test_the_same_values_on_the_form_change_nothing() -> None:
    target = store()
    client, _ = client_for(Role.BILLING)
    client.post(f"/stores/{target.pk}/cost/", cost_form())
    client.post(f"/stores/{target.pk}/cost/", cost_form(effective_from="2026-05-01"))
    assert StoreCost.objects.count() == 1


def test_set_cost_is_billing_and_a_viewer_is_refused() -> None:
    rule = registry.rule_for("ledger:stores:set_cost")
    assert rule is not None and rule.min_role is Role.BILLING and rule.methods == {"POST"}
    target = store()
    viewer, _ = client_for(Role.VIEWER)
    path = f"/stores/{target.pk}/cost/"
    assert viewer.get(path).status_code == 403
    assert viewer.post(path, cost_form()).status_code == 403
    assert 'id="change-cost"' not in page(viewer, target)
    assert not StoreCost.objects.exists()
    admin, _ = client_for(Role.ADMIN)
    assert 'id="change-cost"' in page(admin, target)
    assert admin.post(path, cost_form()).status_code == 302
    assert StoreCost.objects.count() == 1


def test_an_unknown_store_is_404() -> None:
    client, _ = client_for(Role.BILLING)
    assert client.get(f"/stores/{uuid.uuid4()}/cost/").status_code == 404
    assert client.get("/stores/not-a-uuid/cost/").status_code == 404
